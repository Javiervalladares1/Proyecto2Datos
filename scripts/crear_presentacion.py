"""
Genera la presentación de resultados del Proyecto 2.

El contenido proviene de docs/Presentacion_Contenido.md y las figuras de
figures/, ambas producidas por el notebook. Se genera con python-pptx para que
la presentación pueda reconstruirse si cambia el análisis.

Uso (desde la raíz del proyecto):
    python3 scripts/crear_presentacion.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

RAIZ = Path(__file__).resolve().parent.parent
FIGURAS = RAIZ / "figures"
DESTINO = RAIZ / "docs" / "Proyecto2_Presentacion.pptx"

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0F, 0x6E, 0x63)
GRIS = RGBColor(0x44, 0x4A, 0x52)

ANCHO, ALTO = Inches(13.333), Inches(7.5)


def nueva(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # en blanco


def titulo(slide, texto, subtitulo=None):
    caja = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), ANCHO - Inches(1.2), Inches(0.9))
    p = caja.text_frame.paragraphs[0]
    p.text = texto
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = AZUL
    if subtitulo:
        p2 = caja.text_frame.add_paragraph()
        p2.text = subtitulo
        p2.font.size = Pt(15)
        p2.font.color.rgb = GRIS


def vinetas(slide, items, izquierda=0.7, arriba=1.5, ancho=None, tam=17):
    ancho = ancho or (ANCHO - Inches(1.4))
    caja = slide.shapes.add_textbox(Inches(izquierda), Inches(arriba), ancho, Inches(4.4))
    tf = caja.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {item}" if not item.startswith(" ") else item
        p.font.size = Pt(tam)
        p.font.color.rgb = GRIS
        p.space_after = Pt(10)


def mensaje(slide, texto):
    caja = slide.shapes.add_textbox(Inches(0.7), ALTO - Inches(0.78), ANCHO - Inches(1.4), Inches(0.55))
    p = caja.text_frame.paragraphs[0]
    p.text = texto
    p.font.size = Pt(16)
    p.font.italic = True
    p.font.bold = True
    p.font.color.rgb = TEAL


def imagen(slide, nombre, izquierda, arriba, ancho):
    ruta = FIGURAS / nombre
    if ruta.exists():
        slide.shapes.add_picture(str(ruta), Inches(izquierda), Inches(arriba), width=Inches(ancho))


prs = Presentation()
prs.slide_width, prs.slide_height = ANCHO, ALTO

# --- 1. Portada -------------------------------------------------------------
s = nueva(prs)
caja = s.shapes.add_textbox(Inches(0.9), Inches(2.1), ANCHO - Inches(1.8), Inches(3))
tf = caja.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "A Step Ahead of Drought"
p.font.size = Pt(46); p.font.bold = True; p.font.color.rgb = AZUL
for texto, tam, color in [
    ("Pronóstico global de almacenamiento de agua · Reto ITU", 22, TEAL),
    ("Proyecto 2 — Análisis Exploratorio", 20, GRIS),
    ("CC3084 Data Science · Universidad del Valle de Guatemala · Semestre II 2026", 15, GRIS),
    ("Integrantes: ______________________", 14, GRIS),
    ("https://github.com/Javiervalladares1/CC3084-Proyecto2-Drought-EDA", 12, TEAL),
]:
    q = tf.add_paragraph(); q.text = texto
    q.font.size = Pt(tam); q.font.color.rgb = color

# --- 2. Contexto ------------------------------------------------------------
s = nueva(prs)
titulo(s, "Contexto del problema")
vinetas(s, [
    "El TWS (Total Water Storage) es toda el agua sobre y bajo la superficie: "
    "subterránea, humedad del suelo, superficial, nieve y hielo.",
    "Las misiones GRACE y GRACE-FO lo estiman midiendo variaciones del campo "
    "gravitatorio terrestre.",
    "Es el mejor indicador integrado de sequía hidrológica a gran escala.",
], arriba=1.4, ancho=Inches(5.4), tam=16)
imagen(s, "18_mapa_promedio.png", 6.1, 2.0, 6.8)
mensaje(s, "Medimos el agua del planeta pesándolo desde el espacio.")

# --- 3. Problema y objetivos ------------------------------------------------
s = nueva(prs)
titulo(s, "Problema científico y objetivos")
vinetas(s, [
    "El problema: los productos GRACE se publican con 2 a 3 meses de retraso. "
    "Para alerta temprana se necesita el estado actual, y se tiene el de hace un trimestre.",
    "El reto: predecir el TWS del mes siguiente. Métrica de evaluación: RMSE.",
    "Pregunta científica: ¿qué patrones temporales, estacionales y espaciales presenta "
    "el TWS, y qué variables se asocian más con el valor del mes siguiente?",
    "Objetivo: caracterizar esos patrones para sustentar un modelo posterior de "
    "pronóstico a un mes.",
])
mensaje(s, "La información existe, pero llega tarde para decidir.")

# --- 4. Dataset -------------------------------------------------------------
s = nueva(prs)
titulo(s, "Descripción del dataset")
vinetas(s, [
    "2 154 021 observaciones de entrenamiento · 280 961 de prueba · 13 variables",
    "Rejilla global de 1°, con 15 715 celdas, de −55.5° a 83.5° de latitud",
    "Periodo: mayo 2002 – agosto 2015 (138 meses observados)",
    "Variables: TWS_t, SPEI a 1/3/6/12 meses, SOIL_MOISTURE_t y target",
    "Todas estandarizadas y adimensionales: no son centímetros de agua",
])
mensaje(s, "Una fila por mes y por celda del planeta.")

# --- 5. Calidad -------------------------------------------------------------
s = nueva(prs)
titulo(s, "Calidad y limpieza de los datos")
vinetas(s, [
    "Cero nulos y cero duplicados. Pero eso engaña.",
    "El formato es plano: un mes no observado no genera filas, en vez de generar NaN.",
    "Faltan 22 de los 160 meses del calendario (13.8 %), en 2002–2003 y 2011–2014.",
    "Atípicos: 0.25 % en target, 2.58 % en humedad del suelo. No se eliminaron: se "
    "agrupan en años y regiones concretos y son lo que el reto busca anticipar.",
], arriba=1.4, ancho=Inches(5.6), tam=15)
imagen(s, "10_cobertura_temporal.png", 6.4, 2.3, 6.5)
mensaje(s, "isna().sum() daba cero y aun así faltaba el 14 % del periodo.")

# --- 6. Patrones temporales -------------------------------------------------
s = nueva(prs)
titulo(s, "Principales patrones temporales")
vinetas(s, [
    "Tendencia descendente: el promedio anual cae de 0.322 en 2004 a −0.136 en 2015.",
    "La serie no es estacionaria (Dickey-Fuller, p = 0.526). La tendencia explica el "
    "64 % de la varianza; el ciclo anual, solo el 3 %.",
    "No depende de la agregación: promedio simple y ponderado por área correlacionan 0.936.",
], arriba=1.3, tam=15)
imagen(s, "16_media_movil.png", 1.6, 2.85, 10.1)
mensaje(s, "El registro muestra un descenso sostenido, y la serie tiene memoria larga.")

# --- 7. Estacionalidad ------------------------------------------------------
s = nueva(prs)
titulo(s, "La estacionalidad que casi no vemos")
vinetas(s, [
    "A primera vista no hay estacionalidad: solo el 3 % de la varianza.",
    "La causa: los hemisferios tienen ciclos opuestos y se cancelan al promediarse. "
    "El norte promedia positivo los 12 meses; el sur cae a −0.105 en junio.",
    "Además, el 80.3 % de las celdas están en el hemisferio norte. Al separar por "
    "banda de latitud, el ciclo anual reaparece.",
], arriba=1.3, tam=15)
imagen(s, "17_perfil_estacional.png", 1.4, 3.0, 10.5)
mensaje(s, "Promediar el planeta entero borra el fenómeno. Hay que mirar por región.")

# --- 8. Relaciones ----------------------------------------------------------
s = nueva(prs)
titulo(s, "Principales relaciones entre variables")
imagen(s, "19_matriz_correlacion.png", 0.7, 1.35, 5.6)
vinetas(s, [
    "Correlación con target:",
    "     TWS_t (persistencia) . . . 0.803",
    "     SPEI_12_t . . . . . . . . . . 0.380",
    "     SPEI_06_t . . . . . . . . . . 0.368",
    "     SOIL_MOISTURE_t . . . . 0.320",
    "     SPEI_03_t . . . . . . . . . . 0.285",
    "     SPEI_01_t . . . . . . . . . . 0.204",
    "Los SPEI se ordenan exactamente por escala de acumulación.",
    "Relaciones lineales: Pearson y Spearman difieren ≤ 0.024.",
    "Asociación, no causalidad.",
], izquierda=6.7, arriba=1.5, ancho=Inches(6.0), tam=14)
mensaje(s, "El agua tiene inercia, y la sequía de largo plazo importa más que la del mes pasado.")

# --- 9. Hallazgos -----------------------------------------------------------
s = nueva(prs)
titulo(s, "Hallazgos más importantes")
vinetas(s, [
    "Sin nulos ni duplicados, pero faltan 22 meses completos del calendario.",
    "target es el TWS del mes calendario t+1 — verificado celda por celda.",
    "Las variables están estandarizadas; no admiten lectura como volumen físico.",
    "Tendencia descendente y serie no estacionaria en media.",
    "La estacionalidad se cancela globalmente pero existe por región.",
    "Desbalance espacial: 80 % hemisferio norte; medias de −0.222 a 0.327 según banda.",
    "Persistencia dominante (r = 0.803).",
    "Línea base a superar: RMSE 0.5724.",
], tam=15)

# --- 10. Conclusiones y próximos pasos --------------------------------------
s = nueva(prs)
titulo(s, "Conclusiones y próximos pasos")
vinetas(s, [
    "El problema de calidad no son los nulos, es la cobertura del calendario.",
    "El promedio global no es representativo: hay que analizar por región.",
    "La inercia del sistema es la señal más fuerte disponible.",
], arriba=1.35, ancho=Inches(5.7), tam=15)
vinetas(s, [
    "Reportar siempre contra la línea base (RMSE 0.5724).",
    "Construir rezagos respetando los meses ausentes.",
    "Incorporar la dimensión regional y validar por regiones.",
    "Tratar la tendencia explícitamente.",
    "Priorizar SPEI_12_t y SPEI_06_t.",
    "Validación cronológica, nunca aleatoria.",
    "Manejar el enmascaramiento del 66.53 % de TWS_t en test.",
], izquierda=7.0, arriba=1.35, ancho=Inches(5.8), tam=14)

for etiqueta, izq in [("Conclusiones", 0.7), ("Próximos pasos", 7.0)]:
    caja = s.shapes.add_textbox(Inches(izq), Inches(1.05), Inches(4), Inches(0.35))
    p = caja.text_frame.paragraphs[0]
    p.text = etiqueta
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = TEAL

DESTINO.parent.mkdir(parents=True, exist_ok=True)
prs.save(DESTINO)
print(f"Presentación creada: {DESTINO.relative_to(RAIZ)}  ({len(prs.slides.__iter__.__self__._sldIdLst)} diapositivas)")
